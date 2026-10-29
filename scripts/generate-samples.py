from pathlib import Path
import json, hashlib
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'public'/'samples'
data=json.loads((OUT/'northwind-seed.json').read_text(encoding='utf-8'))
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='KTitle',fontName='Helvetica-Bold',fontSize=27,leading=32,textColor=colors.HexColor('#1d352f'),spaceAfter=20))
styles.add(ParagraphStyle(name='KBody',fontName='Helvetica',fontSize=9,leading=14,spaceAfter=10,textColor=colors.HexColor('#435644')))
styles.add(ParagraphStyle(name='KSmall',fontName='Helvetica',fontSize=7,leading=10,textColor=colors.HexColor('#62775c')))
styles.add(ParagraphStyle(name='KH',fontName='Helvetica-Bold',fontSize=15,leading=20,spaceBefore=10,spaceAfter=14,textColor=colors.HexColor('#1d352f')))
def P(s,style='KBody'): return Paragraph(escape(s),styles[style])
def footer(canvas,doc):
    canvas.setFont('Helvetica',7);canvas.setFillColor(colors.HexColor('#66805b'))
    canvas.drawString(42,815,'KESTREL / NORTHWIND STUDIO / Demo - sample data')
    canvas.drawString(42,26,'Copyright (c) 2026 Mariano Campuzano. All rights reserved.')
    canvas.drawRightString(553,26,str(doc.page))
def grid(headers,rows,widths):
    t=Table([[P(x,'KSmall') for x in headers]]+[[P(str(x),'KSmall') for x in r] for r in rows],colWidths=widths,repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf0df')),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.35,colors.HexColor('#dce5d2')),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8)]))
    return t
story=[Spacer(1,50),P('Northwind Studio AI governance evidence pack','KTitle'),P('October 2026 / Fictional sample / Version 1','KH'),P('Demo - sample data'),P('This sample illustrates how an AI inventory, use case review, training register and transparency records can be assembled for a client or auditor. All companies, people, tools and activity are fictional. This document does not certify legal compliance.'),Spacer(1,24),grid(['RECORD','SAMPLE COVERAGE'],[['AI inventory','14 fictional tools'],['Literacy register','23 fictional employees; 18 marked complete'],['Review status','11 reviewed; 3 awaiting review'],['Transparency','Chat, email and voice sample wording'],['Integrity','SHA-256 manifest; public sample is unsigned']],[160,350]),Spacer(1,25),P('Scope and limitations','KH'),P('Findings are illustrative and prewritten. Classification depends on the actual system, intended purpose, affected people, organization role and applicable law. Course completion and employee acknowledgment support a training record; they are not official regulatory certificates.'),P('The editable DOCX and PDF represent the same sample content. Editing either file changes its hash. Obtain the manifest from the same sample bundle and compare the file bytes. This public manifest is not digitally signed.'),PageBreak(),P('AI inventory and sample review status','KH')]
toolrows=[]
for tool in data['tools']:
    toolrows.append([tool['name'],tool['purpose'],tool['owner'],tool['risk']+' / '+tool['status']])
story += [grid(['TOOL','USE CASE','OWNER','SAMPLE FINDING'],toolrows,[88,174,108,140]),PageBreak(),P('AI literacy register','KH'),P('These are fictitious historical records for demonstration. No real employees completed these courses. The sample course is Working responsibly with AI, version 1. A real record would link an authenticated assignment, course version, quiz result and acknowledgment.'),grid(['EMPLOYEE','TEAM','STATUS','SAMPLE DATE'],[[e['name'],e['team'],'Completed' if e['completed'] else 'In progress' if e['progress'] else 'Not started',e['date'] or 'Not completed'] for e in data['employees']],[145,140,110,115]),PageBreak(),P('Transparency notices and outstanding actions','KH')]
notices=[('Chat','You are chatting with an AI assistant. Its responses may contain errors. Ask to speak with a person whenever you need help.'),('Email','This message was prepared with AI assistance and reviewed by our team. Contact us if you need clarification.'),('Voice','You are listening to an AI-generated voice. You can request assistance from a member of our team.')]
for title,body in notices: story += [P(title,'KH'),P(body)]
story += [P('Applicability and placement','KH'),P('These are example texts, not a legal conclusion that every channel requires a notice. Assess provider and deployer responsibilities, content type, exceptions, timing and accessibility. A textual notice does not replace technical marking where applicable.'),P('Outstanding actions','KH'),P('1. Talent Lens: specialist review of recruitment ranking and applicable high-risk conditions.'),P('2. Meetnote: review meeting data, vendor terms, retention and participant communication.'),P('3. Voice Loom: confirm synthetic voice disclosure and its presentation.'),P('4. Five fictional employees: complete their assigned training.'),P('Sources and versioning','KH'),P('EU AI Act: https://eur-lex.europa.eu/eli/reg/2024/1689/oj/eng'),P('GDPR: https://eur-lex.europa.eu/eli/reg/2016/679/oj/eng'),P('Legal references must be checked against the applicable consolidated text and guidance before production use. Sample generated for the fictional October 2026 snapshot.')]
SimpleDocTemplate(str(OUT/'northwind-evidence.pdf'),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=52,bottomMargin=45,title='Northwind Studio AI governance evidence pack',author='Kestrel').build(story,onFirstPage=footer,onLaterPages=footer)

doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(8)
for style in ['Title','Heading 1','Heading 2']:doc.styles[style].font.color.rgb=RGBColor.from_string('1D352F')
doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.add_paragraph('Northwind Studio AI governance evidence pack','Title');doc.add_paragraph('October 2026 / Fictional sample / Version 1');doc.add_paragraph('Demo - sample data')
doc.add_paragraph('This sample illustrates an AI inventory, use case review, training register and transparency records for clients and auditors. All companies, people, tools and activity are fictional. It does not certify legal compliance.')
doc.add_heading('Scope and limitations',level=1);doc.add_paragraph('Findings are illustrative and prewritten. Applicability depends on system, use case, affected people, organizational role and law. Training records support evidence of participation, not an official regulatory certificate. The public integrity manifest is unsigned; editing either file changes its hash.')
doc.add_paragraph('14 tools / 23 employees / 18 completed / 11 reviewed / 3 awaiting review')
doc.add_page_break();doc.add_heading('AI inventory and sample review status',level=1)
table=doc.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
for c,v in zip(table.rows[0].cells,['Tool','Use case','Owner','Finding and status']):c.text=v
for row in toolrows:
    for c,v in zip(table.add_row().cells,row):c.text=v
doc.add_page_break();doc.add_heading('AI literacy register',level=1);doc.add_paragraph('Fictional historical records. No real employees completed these courses. Course: Working responsibly with AI version 1. Production records link authenticated assignments, quizzes and acknowledgments.')
table=doc.add_table(rows=1,cols=4);table.style='Light Shading Accent 1'
for c,v in zip(table.rows[0].cells,['Employee','Team','Status','Sample date']):c.text=v
for e in data['employees']:
    values=[e['name'],e['team'],'Completed' if e['completed'] else 'In progress' if e['progress'] else 'Not started',e['date'] or 'Not completed']
    for c,v in zip(table.add_row().cells,values):c.text=v
doc.add_page_break();doc.add_heading('Transparency notices and outstanding actions',level=1)
for title,body in notices:doc.add_heading(title,level=2);doc.add_paragraph(body)
doc.add_paragraph('Assess applicability, role, content, exceptions, timing and accessibility before use. Text notices do not replace technical marking. AI-assisted email does not automatically require a disclosure.')
doc.add_heading('Outstanding actions',level=2)
for s in ['Review recruitment ranking for Talent Lens.','Review meeting data and vendor terms for Meetnote.','Confirm disclosure placement for Voice Loom.','Complete training for five fictional employees.']:doc.add_paragraph(s,style='List Bullet')
doc.add_heading('Sources',level=2);doc.add_paragraph('EU AI Act: https://eur-lex.europa.eu/eli/reg/2024/1689/oj/eng');doc.add_paragraph('GDPR: https://eur-lex.europa.eu/eli/reg/2016/679/oj/eng');doc.add_paragraph('Check applicable consolidated text and guidance before production use.')
for section in doc.sections:section.footer.paragraphs[0].text='Demo - sample data | Copyright (c) 2026 Mariano Campuzano. All rights reserved.'
doc.save(OUT/'northwind-evidence.docx')
SimpleDocTemplate(str(OUT/'sample-certificate.pdf'),pagesize=A4,rightMargin=50,leftMargin=50,topMargin=70,bottomMargin=50).build([Spacer(1,70),P('AI literacy completion certificate','KTitle'),P('Demo - sample data'),Spacer(1,25),P('Fictional employee example','KH'),P('Working responsibly with AI / Course version 1'),P('Illustrative completion date: October 21, 2026'),P('Certificate reference: SAMPLE-NW-2026-001'),Spacer(1,40),P('This sample is a layout example. It is not issued to an actual person and does not certify legal compliance. In the private product, certificates are linked to authenticated quiz results and acknowledgments.')],onFirstPage=footer)
files={name:hashlib.sha256((OUT/name).read_bytes()).hexdigest() for name in ['northwind-evidence.pdf','northwind-evidence.docx','sample-certificate.pdf']}
(OUT/'manifest.json').write_text(json.dumps({'label':'Demo - sample data','signed':False,'algorithm':'SHA-256','files':files},indent=2)+'\n',encoding='utf-8')
print('Generated sample PDF, DOCX, certificate and unsigned SHA-256 manifest.')
